#!/usr/bin/env python3
"""Audit the evidence needed for the publication-facing GitHub release."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from zipfile import ZipFile

import duckdb
import pandas as pd
from docx import Document
from lxml import etree


RELEASE_VERSION = "1.2.1"
RELEASE_DATE = "2026-10-11"
EXPECTED_COHORTS = {
    "TCGA-BRCA": (1064, 150, 1064),
    "GSE96058": (3273, 336, 3273),
    "METABRIC": (1980, 1143, 1979),
    "GSE20685": (327, 83, 327),
}
EXPECTED_DEVIATIONS = {
    "Cohorts",
    "Cohort inclusion",
    "Model inputs",
    "Pathway definitions",
    "Candidate models",
    "Internal cross-validation",
    "Hyperparameter search",
    "PAM50 comparator",
    "Oncotype DX comparison",
    "Primary endpoint: training cohort",
    "Secondary (a) ML vs Cox >= +0.02",
    "Secondary (b) immune pathway dominant in TNBC and HER2+",
    "Secondary (c) combined ML + PAM50 + clinical IDI > 0.05",
    "Secondary (d) non-inferiority to PAM50, margin -0.02",
    "Post-hoc analyses",
    "Registration document",
}
EXPLICIT_CHRONOLOGY = "Analyses first performed after registration were"
AI_SCOPE = "used only for grammar, language, and formatting checks"
PROHIBITED_AI_ATTRIBUTION = (
    "claude",
    "anthropic",
    "gemini",
    "conceptual layout",
    "graphic elements of figure 1",
    "co-authored-by",
)
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _require(condition: bool, message: str, issues: list[str]) -> None:
    if not condition:
        issues.append(message)


def _version_checks(root: Path, issues: list[str]) -> None:
    version_path = root / "VERSION"
    _require(version_path.exists(), "VERSION is missing", issues)
    if version_path.exists():
        _require(_read(version_path).strip() == RELEASE_VERSION, "VERSION is not 1.2.1", issues)

    citation = _read(root / "CITATION.cff")
    _require(f'version: "{RELEASE_VERSION}"' in citation, "CITATION.cff version is not 1.2.1", issues)
    _require(
        f'date-released: "{RELEASE_DATE}"' in citation,
        "CITATION.cff release date is not 2026-10-11",
        issues,
    )

    readme = _read(root / "README.md")
    for required in (
        "6,644",
        "6,643",
        "DeepSurv",
        "+0.020",
        "v1.2.1",
        "Analyses first performed after registration",
        "Supplementary Information",
    ):
        _require(required in readme, f"README is missing publication text: {required}", issues)
    for stale in ("4,532 samples", "delta = +0.0144", "official genefu PAM50-ROR"):
        _require(stale not in readme, f"README still contains superseded text: {stale}", issues)

    workflow = _read(root / ".github" / "workflows" / "ci.yml")
    _require(
        "python scripts/check_publication_readiness.py" in workflow,
        "CI does not run the publication-readiness gate",
        issues,
    )


def _cohort_checks(root: Path, issues: list[str]) -> None:
    flow = json.loads(_read(root / "results" / "Figure1_cohort_flow_counts.json"))
    analysed = sum(int(row["analysed"]) for row in flow.values())
    evaluable = sum(int(row["evaluable"]) for row in flow.values())
    _require(analysed == 6644, f"analysed total is {analysed}, expected 6644", issues)
    _require(evaluable == 6643, f"PAM50-evaluable total is {evaluable}, expected 6643", issues)

    crosswalk = pd.read_csv(root / "results" / "Table_1_denominator_crosswalk.csv").set_index("cohort")
    for cohort, (cohort_n, events, cohort_evaluable) in EXPECTED_COHORTS.items():
        _require(cohort in flow, f"cohort flow is missing {cohort}", issues)
        _require(cohort in crosswalk.index, f"denominator crosswalk is missing {cohort}", issues)
        if cohort not in flow or cohort not in crosswalk.index:
            continue
        _require(int(flow[cohort]["analysed"]) == cohort_n, f"{cohort} analysed n is wrong", issues)
        _require(int(flow[cohort]["analysed_events"]) == events, f"{cohort} event count is wrong", issues)
        _require(
            int(flow[cohort]["evaluable"]) == cohort_evaluable,
            f"{cohort} evaluable n is wrong",
            issues,
        )
        _require(
            int(crosswalk.loc[cohort, "patient_characteristics_final_n"]) == cohort_n,
            f"{cohort} crosswalk n does not match cohort flow",
            issues,
        )


def _registration_checks(root: Path, issues: list[str]) -> None:
    results = root / "results"
    s18 = pd.read_csv(results / "Table_S18_registered_endpoints.csv")
    _require(len(s18) == 8, f"Table S18 has {len(s18)} rows, expected 8", issues)
    oncotype_rows = s18["endpoint"].str.contains("Oncotype DX", case=False, na=False)
    _require(int(oncotype_rows.sum()) == 2, "Table S18 must contain both Oncotype DX dispositions", issues)
    _require(
        s18["registered_criterion"].notna().all(),
        "Table S18 has a registered endpoint without a criterion",
        issues,
    )

    s20 = pd.read_csv(results / "Table_S20_registration_deviations.csv")
    _require(len(s20) == 16, f"Table S20 has {len(s20)} rows, expected 16", issues)
    actual_deviations = set(s20["item"].astype(str))
    missing = sorted(EXPECTED_DEVIATIONS - actual_deviations)
    _require(not missing, f"Table S20 is missing deviations: {', '.join(missing)}", issues)
    _require(
        not s20.astype(str).apply(lambda column: column.str.contains("registration\\.\\.")).any().any(),
        "Table S20 contains a duplicated period in the registration chronology",
        issues,
    )

    for name in (
        "Table_S19_locked_model_full_cohorts.csv",
        "Table_S21_age_adjusted_comparison.csv",
    ):
        _require((results / name).exists(), f"required supplementary source is missing: {name}", issues)


def _stage_checks(root: Path, issues: list[str]) -> dict[str, int]:
    sys.path.insert(0, str(root))
    from src.clinical import stage_to_ordinal

    database = root / "data" / "processed" / "unified_cohorts.duckdb"
    connection = duckdb.connect(str(database), read_only=True)
    try:
        stages = connection.execute("select cohort, stage_tnm from samples").fetchdf()
    finally:
        connection.close()
    stages["stage_ordinal"] = stages["stage_tnm"].map(stage_to_ordinal)
    counts = {
        str(cohort): int(group["stage_ordinal"].notna().sum())
        for cohort, group in stages.groupby("cohort")
    }
    _require(counts.get("TCGA-BRCA", 0) > 0, "TCGA-BRCA has no parsed stage", issues)
    _require(counts.get("METABRIC", 0) > 0, "METABRIC has no parsed stage", issues)
    _require(counts.get("GSE20685", 0) > 0, "GSE20685 has no parsed stage", issues)
    _require(counts.get("GSE96058", -1) == 0, "GSE96058 unexpectedly has parsed stage", issues)
    return counts


def _supplement_checks(path: Path, issues: list[str]) -> None:
    _require(path.exists(), f"supplement does not exist: {path}", issues)
    if not path.exists():
        return
    document = Document(path)
    _require(len(document.tables) == 4, "supplement must embed exactly Tables S18-S21", issues)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    for label in ("Table S18", "Table S19", "Table S20", "Table S21"):
        _require(label in text, f"supplement is missing label {label}", issues)
    oncotype_dispositions = sum(
        "oncotype dx" in row.cells[0].text.lower()
        for row in document.tables[0].rows[1:]
    )
    _require(
        oncotype_dispositions == 2,
        "supplement must contain both Oncotype DX dispositions",
        issues,
    )
    _require(
        "GSE96058 releases none" in table_text,
        "supplement does not state that GSE96058 releases no stage",
        issues,
    )
    _require(
        EXPLICIT_CHRONOLOGY in text,
        "supplement does not identify the analyses first performed after registration",
        issues,
    )
    _require("release v1.2.1" in text, "supplement does not cite release v1.2.1", issues)
    _require("v1.2.0" not in text, "supplement still cites superseded release v1.2.0", issues)


def _document_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + cells)


def _clean_docx_checks(paths: list[Path], issues: list[str]) -> None:
    for path in paths:
        _require(path.exists(), f"document does not exist: {path}", issues)
        if not path.exists():
            continue
        text = _document_text(path).lower()
        for phrase in PROHIBITED_AI_ATTRIBUTION:
            _require(phrase not in text, f"{path.name} contains prohibited AI attribution: {phrase}", issues)
        with ZipFile(path) as archive:
            names = set(archive.namelist())
            _require(
                not any(name.startswith("word/comments") for name in names),
                f"{path.name} contains Word comments",
                issues,
            )
            for name in names:
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                root = etree.fromstring(archive.read(name))
                tracked = root.xpath(".//w:ins | .//w:del", namespaces={"w": WORD_NS})
                _require(not tracked, f"{path.name} contains tracked changes in {name}", issues)


def _manuscript_checks(path: Path, issues: list[str]) -> None:
    _require(path.exists(), f"manuscript does not exist: {path}", issues)
    if not path.exists():
        return
    text = _document_text(path)
    lower = text.lower()
    _require(EXPLICIT_CHRONOLOGY in text, "manuscript chronology is not explicit", issues)
    _require(AI_SCOPE in lower, "manuscript does not limit ChatGPT to language checks", issues)
    _require(lower.count("chatgpt (openai)") == 1, "manuscript AI disclosure must appear once", issues)
    _require("release v1.2.1" in text, "manuscript does not cite release v1.2.1", issues)
    _require("v1.2.0" not in text, "manuscript still cites superseded release v1.2.0", issues)


def _response_checks(path: Path, issues: list[str]) -> None:
    _require(path.exists(), f"response does not exist: {path}", issues)
    if not path.exists():
        return
    text = _document_text(path)
    lower = text.lower()
    for filename in (
        "Manuscript_prepublication_revision_v1.2.1.docx",
        "BMC_Supplement_prepublication_revision_v1.2.1.docx",
        "Response_to_prepublication_checks_v1.2.1.docx",
    ):
        _require(filename in text, f"response does not list {filename}", issues)
    _require(EXPLICIT_CHRONOLOGY.lower() in lower, "response chronology is not explicit", issues)
    _require(AI_SCOPE in lower, "response does not limit ChatGPT to language checks", issues)

    with ZipFile(path) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
        colors = {
            element.get(f"{{{WORD_NS}}}val", "").upper()
            for element in styles.xpath(".//w:color", namespaces={"w": WORD_NS})
        }
        fills = {
            element.get(f"{{{WORD_NS}}}fill", "").upper()
            for element in styles.xpath(".//w:shd", namespaces={"w": WORD_NS})
        }
        border_colors = {
            element.get(f"{{{WORD_NS}}}color", "").upper()
            for element in styles.xpath(".//w:pBdr/*", namespaces={"w": WORD_NS})
        }
        _require(colors <= {"000000", "AUTO"}, f"response styles contain non-black text: {sorted(colors)}", issues)
        _require(fills <= {"FFFFFF", "AUTO", "NONE"}, f"response styles contain non-white fills: {sorted(fills)}", issues)
        _require(
            border_colors <= {"000000", "AUTO"},
            f"response styles contain non-black paragraph borders: {sorted(border_colors)}",
            issues,
        )


def audit_repository(
    root: Path,
    manuscript: Path | None = None,
    supplement: Path | None = None,
    response: Path | None = None,
) -> list[str]:
    """Return every publication-readiness failure found under ``root``."""

    issues: list[str] = []
    _version_checks(root, issues)
    _cohort_checks(root, issues)
    _registration_checks(root, issues)
    stage_counts = _stage_checks(root, issues)
    if manuscript is not None:
        _manuscript_checks(manuscript, issues)
    if supplement is not None:
        _supplement_checks(supplement, issues)
    if response is not None:
        _response_checks(response, issues)
    documents = [path for path in (manuscript, supplement, response) if path is not None]
    if documents:
        _clean_docx_checks(documents, issues)
    audit_repository.stage_counts = stage_counts  # type: ignore[attr-defined]
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--manuscript", type=Path)
    parser.add_argument("--supplement", type=Path)
    parser.add_argument("--response", type=Path)
    arguments = parser.parse_args()

    root = arguments.root.resolve()
    manuscript = arguments.manuscript.resolve() if arguments.manuscript else None
    supplement = arguments.supplement.resolve() if arguments.supplement else None
    response = arguments.response.resolve() if arguments.response else None
    issues = audit_repository(root, manuscript, supplement, response)
    if issues:
        for issue in issues:
            print(f"FAIL {issue}")
        return 1

    counts = audit_repository.stage_counts  # type: ignore[attr-defined]
    print(f"PASS version: {RELEASE_VERSION}")
    print("PASS cohorts: analysed=6644; PAM50-evaluable=6643")
    print("PASS registered endpoints: 8 rows; Oncotype DX dispositions=2")
    print("PASS registration deviations: 16 rows")
    print(
        "PASS stage availability: "
        + "; ".join(f"{cohort}={counts[cohort]}" for cohort in sorted(counts))
    )
    if manuscript is not None:
        print("PASS manuscript: chronology explicit; ChatGPT limited to language checks")
    if supplement is not None:
        print("PASS supplement: Tables S18-S21 embedded; Oncotype DX dispositions=2")
    if response is not None:
        print("PASS response: versioned files listed; black text on white background")
    if any(path is not None for path in (manuscript, supplement, response)):
        print("PASS clean DOCX: no tracked changes, comments, or prohibited AI attribution")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
